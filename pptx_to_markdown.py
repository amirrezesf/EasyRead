"""Extract a .pptx deck into a single Markdown file that preserves formatting
cues (bold / underline / font color / highlight) and extracts embedded images
with Markdown image references in reading order.

Usage:
    python pptx_to_markdown.py input.pptx output.md [--images-dir images]

Design notes / limitations (read before wiring this into your app):

1. "Reading order" on a slide is a heuristic. PPTX shapes have no native
   document order — only a z-order (creation order) and (x, y) coordinates.
   This script sorts shapes by a column-aware heuristic: shapes are clustered
   into vertical columns by their horizontal centers, then sorted top-to-bottom
   within each column, columns processed left-to-right. This handles simple
   multi-column layouts better than naive (top, left) but is still a heuristic.
   Spot check a few slides before trusting it on your whole deck.

2. "Important" formatting is inferred from bold / underline / font color /
   highlight. This is only as reliable as the professor's own habits — if
   color/underline was used decoratively (or inconsistently) rather than to
   mark importance, this will produce false positives/negatives. Skim the
   Markdown output for a slide or two to sanity check before trusting it across
   the whole deck.

3. Images are extracted as files only — this script does NOT caption or
   interpret them. The Markdown image reference uses a generic alt text. If you
   want the downstream LLM to actually understand a chart/diagram, you need
   a separate vision-model pass per image that writes a real caption into
   the alt text before you hand the Markdown to a text-only model.
"""

import argparse
import os
import re
from pathlib import Path
from typing import Callable, Optional, Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from convert_emf_and_wmf import vector_to_jpg


# ---------- formatting detection ----------

def _run_highlight_rgb(run) -> Optional[str]:
    """python-pptx has no public API for <a:highlight> (text highlight color)."""
    r = run._r
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    hl = r.find("a:rPr/a:highlight/a:srgbClr", ns)
    if hl is not None:
        return hl.get("val")
    return None


def _run_font_rgb(run) -> Optional[str]:
    try:
        color = run.font.color
        if color is not None and color.type is not None and color.rgb is not None:
            return str(color.rgb)
    except Exception:
        pass
    return None


def _run_is_emphasized(run) -> tuple[bool, dict[str, Optional[bool | str]]]:
    bold = bool(run.font.bold)
    underline = bool(run.font.underline)
    color = _run_font_rgb(run)
    highlight = _run_highlight_rgb(run)
    color_is_signal = color is not None and color.upper() not in ("000000", "FFFFFF")
    return bold or underline or color_is_signal or bool(highlight), {
        "bold": bold,
        "underline": underline,
        "color": color if color_is_signal else None,
        "highlight": highlight,
    }


def _escape_md(text: str) -> str:
    return (
        text.replace("\\", "\\\\")
        .replace("*", "\\*")
        .replace("_", "\\_")
        .replace("`", "\\`")
        .replace("[", "\\[")
        .replace("]", "\\]")
    )


def _run_to_markdown(run) -> str:
    text = _escape_md(run.text)
    if not text:
        return ""
    _, flags = _run_is_emphasized(run)
    if flags["highlight"]:
        text = f"**=={text}==**"  # bold + highlight marker for visibility
    elif flags["color"]:
        text = f"**{text}**"  # bold for colored text (color not representable in standard MD)
    if flags["bold"]:
        text = f"**{text}**"
    if flags["underline"]:
        text = f"__{text}__"  # underline as double underscore
    return text


def _paragraph_to_markdown(paragraph) -> Optional[tuple[int, str]]:
    runs_md = "".join(_run_to_markdown(r) for r in paragraph.runs)
    if not runs_md.strip():
        return None
    level = paragraph.level or 0
    return level, runs_md


def _text_frame_to_markdown(text_frame) -> str:
    lines = []
    for p in text_frame.paragraphs:
        result = _paragraph_to_markdown(p)
        if result is None:
            continue
        level, md = result
        indent = "  " * level
        lines.append(f"{indent}- {md}")
    return "\n".join(lines)


def _cell_plain(cell) -> str:
    parts = []
    for p in cell.text_frame.paragraphs:
        result = _paragraph_to_markdown(p)
        if result is not None:
            parts.append(result[1].replace("|", "\\|"))
    return " ".join(parts).strip()


def _table_to_markdown(table) -> str:
    rows = []
    for row in table.rows:
        cells = [_cell_plain(cell) for cell in row.cells]
        rows.append(cells)
    if not rows:
        return ""
    col_count = max(len(r) for r in rows)
    for row in rows:
        while len(row) < col_count:
            row.append("")
    header = rows[0]
    body = rows[1:] if len(rows) > 1 else []
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    if not body:
        lines.append("| " + " | ".join("" for _ in header) + " |")
    else:
        for row in body:
            lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _picture_to_markdown(
    shape,
    slide_idx: int,
    img_counter: int,
    images_dir: str,
    md_out_dir: str,
) -> str:
    image = shape.image
    ext = image.ext.lower()
    filename = f"slide{slide_idx}_img{img_counter}.{ext}"
    filepath = os.path.join(images_dir, filename)
    with open(filepath, "wb") as f:
        f.write(image.blob)
    if ext in {"emf", "wmf"}:
        converted_path = os.path.splitext(filepath)[0] + ".jpg"
        vector_to_jpg(filepath, converted_path)
        os.remove(filepath)
        filepath = converted_path
    rel_path = os.path.relpath(filepath, md_out_dir).replace("\\", "/")
    alt = f"Slide {slide_idx} image {img_counter}"
    return f"![{alt}]({rel_path})"


# ---------- column-aware reading order ----------

def _shape_center_x(shape) -> float:
    """Return the horizontal center of a shape in EMU."""
    left = shape.left if shape.left is not None else 0
    width = shape.width if shape.width is not None else 0
    return left + width / 2


def _shape_top(shape) -> int:
    """Return the top position of a shape in EMU."""
    return shape.top if shape.top is not None else 0


def _cluster_into_columns(
    shapes: list,
    column_threshold_emu: int = 914400,  # ~1 inch in EMU
) -> list[list]:
    """
    Cluster shapes into vertical columns by horizontal centers.
    
    Shapes whose horizontal centers are within `column_threshold_emu` of each other
    are placed in the same column. Columns are then ordered left-to-right by their
    median center x. Within each column, shapes are sorted top-to-bottom.
    
    This handles simple multi-column layouts (e.g., two columns of bullets) much
    better than a pure (top, left) sort, which would interleave the columns.
    """
    if not shapes:
        return []
    
    # Compute center x for each shape
    shapes_with_center = [(s, _shape_center_x(s)) for s in shapes]
    
    # Sort by center x to form initial column groups
    shapes_with_center.sort(key=lambda x: x[1])
    
    columns = []
    current_column = []
    current_column_centers = []
    
    for shape, cx in shapes_with_center:
        if not current_column:
            current_column.append(shape)
            current_column_centers.append(cx)
        else:
            # Check if this shape belongs to the current column
            median_cx = sorted(current_column_centers)[len(current_column_centers) // 2]
            if abs(cx - median_cx) <= column_threshold_emu:
                current_column.append(shape)
                current_column_centers.append(cx)
            else:
                # Start new column
                columns.append(current_column)
                current_column = [shape]
                current_column_centers = [cx]
    
    if current_column:
        columns.append(current_column)
    
    # Sort shapes within each column by top position
    for column in columns:
        column.sort(key=_shape_top)
    
    # Columns are already left-to-right by construction
    return columns


def _get_ordered_shapes(slide) -> list:
    """
    Return shapes in reading order using column-aware clustering.
    """
    shapes = list(slide.shapes)
    columns = _cluster_into_columns(shapes)
    ordered = []
    for column in columns:
        ordered.extend(column)
    return ordered


# ---------- shape collection ----------

def _collect_shape_markdown(
    shape,
    slide_idx: int,
    img_counter: int,
    images_dir: str,
    md_out_dir: str,
    pieces: list[str],
) -> int:
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        img_counter += 1
        pieces.append(
            _picture_to_markdown(shape, slide_idx, img_counter, images_dir, md_out_dir)
        )
        return img_counter

    if getattr(shape, "has_table", False) and shape.has_table:
        md = _table_to_markdown(shape.table)
        if md:
            pieces.append(md)
        return img_counter

    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        md = _text_frame_to_markdown(shape.text_frame)
        if md:
            pieces.append(md)
        return img_counter

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in _get_ordered_shapes(type('obj', (object,), {'shapes': list(shape.shapes)})):
            if sub.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img_counter += 1
                pieces.append(
                    _picture_to_markdown(sub, slide_idx, img_counter, images_dir, md_out_dir)
                )
            elif getattr(sub, "has_table", False) and sub.has_table:
                md = _table_to_markdown(sub.table)
                if md:
                    pieces.append(md)
            elif getattr(sub, "has_text_frame", False) and sub.has_text_frame:
                md = _text_frame_to_markdown(sub.text_frame)
                if md:
                    pieces.append(md)
        return img_counter

    return img_counter


def _slide_to_markdown(
    slide,
    slide_idx: int,
    images_dir: str,
    md_out_dir: str,
) -> str:
    pieces = [f"## Slide {slide_idx}"]
    img_counter = 0
    shapes = _get_ordered_shapes(slide)
    for shape in shapes:
        img_counter = _collect_shape_markdown(
            shape, slide_idx, img_counter, images_dir, md_out_dir, pieces
        )
    return "\n\n".join(pieces)


def convert(
    pptx_path: str | Path,
    md_path: str | Path,
    images_dir: Optional[str | Path] = None,
    log: Callable[[str], Any] = print,
) -> tuple[Path, Path]:
    prs = Presentation(pptx_path)
    md_path = Path(md_path)
    md_out_dir = md_path.parent or Path(".")
    if images_dir is None:
        images_dir = md_out_dir / "images"
    else:
        images_dir = Path(images_dir)
    os.makedirs(md_out_dir, exist_ok=True)
    os.makedirs(images_dir, exist_ok=True)

    slides_md = []
    for i, slide in enumerate(prs.slides, start=1):
        slides_md.append(_slide_to_markdown(slide, i, str(images_dir), str(md_out_dir)))
        log(f"Converted slide {i}/{len(prs.slides)}")

    title = md_path.stem
    doc = f"# {title}\n\n" + "\n\n---\n\n".join(slides_md) + "\n"
    # Collapse excessive blank lines
    doc = re.sub(r"\n{3,}", "\n\n", doc)

    with open(md_path, "w", encoding="utf-8") as f:
        f.write(doc)

    return md_path, images_dir


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pptx_path", type=Path)
    ap.add_argument("md_path", type=Path)
    ap.add_argument("--images-dir", type=Path, default=None, help="Defaults to <md_dir>/images")
    args = ap.parse_args()

    md_path, images_dir = convert(args.pptx_path, args.md_path, args.images_dir)
    print(f"Wrote {md_path}")
    print(f"Images in {images_dir}")


if __name__ == "__main__":
    main()
