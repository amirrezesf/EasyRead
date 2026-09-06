from pptx import Presentation

def extract_ppt_text(file_path):
    """
    Extracts all text from a PowerPoint (.pptx) file.

    Args:
        file_path (str): Path to the PowerPoint file.

    Returns:
        str: All text found in the presentation.
    """
    prs = Presentation(file_path)
    text = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        text.append(f"--- Slide {slide_num} ---")

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text.append(shape.text)

    return "\n".join(text)


# Example usage
if __name__ == "__main__":
    file_path = "test.pptx"
    extracted_text = extract_ppt_text(file_path)
    with open("extracted.txt", "w", encoding="utf-8") as f:
        f.write(extracted_text)