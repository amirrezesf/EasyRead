"""Extract a .pptx deck into a single HTML file that preserves formatting cues
(bold / underline / font color / highlight) as semantic emphasis markup, and
extracts embedded images to a folder with <img> references placed at their
approximate position in reading order.

Usage:
    python pptx_to_html.py input.pptx output.html [--images-dir images]

Design notes / limitations (read before wiring this into your app):

1. "Reading order" on a slide is a heuristic. PPTX shapes have no native
   document order — only a z-order (creation order) and (x, y) coordinates.
   This script sorts shapes by a column-aware heuristic: shapes are clustered
   into vertical columns by their horizontal centers, then sorted top-to-bottom
   within each column, columns processed left-to-right. This handles simple
   multi-column layouts better than naive (top, left) but is still a heuristic.
   Spot check a few slides before trusting it on your whole deck.

2. "Important" formatting is inferred from bold / underline / font color /
   highlight. This is only as reliable as the professor's own habits — if
   color/underline was used decoratively (or inconsistently) rather than to
   mark importance, this will produce false positives/negatives. Skim the
   HTML output for a slide or two to sanity check before trusting it across
   the whole deck.

3. Images are extracted as files only — this script does NOT caption or
   interpret them. The <img> tag's alt text is a generic placeholder. If you
   want the downstream LLM to actually understand a chart/diagram, you need
   a separate vision-model pass per image that writes a real caption into
   the alt text or a following <p class="img-caption"> tag before you hand
   the HTML to a text-only model. A text-only LLM reading this HTML sees the
   filename, nothing else.

4. Slide notes (speaker notes) are not extracted. Add `slide.notes_slide`
   handling if you need them.
"""

import argparse
import os
import sys
from html import escape
from pathlib import Path
from typing import Callable, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from convert_emf_and_wmf import vector_to_jpg


# ---------- formatting detection ----------

def _run_highlight_rgb(run) -> Optional[str]:
    """python-pptx has no public API for <a:highlight> (text highlight color).
    Reach into the underlying XML for it."""
    r = run._r
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    hl = r.find("a:rPr/a:highlight/a:srgbClr", ns)
    if hl is not None:
        return hl.get("val")
    return None


def _run_font_rgb(run) -> Optional[str]:
    try:
        color = run.font.color
        if color is not None and color.type is not None and color.rgb is not None:
            return str(color.rgb)
    except Exception:
        pass
    return None


def _run_is_emphasized(run) -> tuple[bool, dict[str, Optional[bool | str]]]:
    """Bold, underline, an explicit font color, or a highlight all count as
    an 'this looks important' signal."""
    bold = bool(run.font.bold)
    underline = bool(run.font.underline)
    color = _run_font_rgb(run)
    highlight = _run_highlight_rgb(run)
    # Ignore plain black/white font colors — decks often set those explicitly
    # without meaning to flag anything.
    color_is_signal = color is not None and color.upper() not in ("000000", "FFFFFF")
    return bold or underline or color_is_signal or bool(highlight), {
        "bold": bold,
        "underline": underline,
        "color": color if color_is_signal else None,
        "highlight": highlight,
    }


def _run_to_html(run) -> str:
    text = escape(run.text)
    if not text:
        return ""
    _, flags = _run_is_emphasized(run)
    if flags["highlight"]:
        text = f'<mark style="background-color:#{flags["highlight"]}">{text}</mark>'
    elif flags["color"]:
        text = f'<span style="color:#{flags["color"]}">{text}</span>'
    if flags["bold"]:
        text = f"<strong>{text}</strong>"
    if flags["underline"]:
        text = f"<u>{text}</u>"
    return text


def _paragraph_to_html(paragraph) -> Optional[tuple[int, str]]:
    runs_html = "".join(_run_to_html(r) for r in paragraph.runs)
    if not runs_html.strip():
        return None
    level = paragraph.level or 0
    return level, runs_html


# ---------- shape handlers ----------

def _text_frame_to_html(text_frame) -> str:
    parts = []
    list_stack = []  # track open <ul> nesting by level

    def close_to(level):
        while list_stack and list_stack[-1] >= level:
            parts.append("</ul>")
            list_stack.pop()

    any_list = False
    for p in text_frame.paragraphs:
        result = _paragraph_to_html(p)
        if result is None:
            continue
        level, html = result
        close_to(level) if list_stack and list_stack[-1] > level else None
        if not list_stack or list_stack[-1] < level:
            parts.append("<ul>")
            list_stack.append(level)
            any_list = True
        elif list_stack[-1] > level:
            close_to(level)
            if not list_stack:
                parts.append("<ul>")
                list_stack.append(level)
        parts.append(f"<li>{html}</li>")
    while list_stack:
        parts.append("</ul>")
        list_stack.pop()
    return "".join(parts) if any_list else ""


def _table_to_html(table) -> str:
    rows_html = []
    for row in table.rows:
        cells_html = []
        for cell in row.cells:
            cell_html = (
                _text_frame_to_html(cell.text_frame)
                .replace("<ul>", "")
                .replace("</ul>", "")
                .replace("<li>", "")
                .replace("</li>", " ")
            )
            cells_html.append(f"<td>{cell_html.strip()}</td>")
        rows_html.append(f"<tr>{''.join(cells_html)}</tr>")
    return f"<table border='1' cellspacing='0' cellpadding='4'>{''.join(rows_html)}</table>"


def _picture_to_html(
    shape,
    slide_idx: int,
    img_counter: int,
    images_dir: str,
    html_out_dir: str,
) -> str:
    image = shape.image
    ext = image.ext.lower()
    filename = f"slide{slide_idx}_img{img_counter}.{ext}"
    filepath = os.path.join(images_dir, filename)
    with open(filepath, "wb") as f:
        f.write(image.blob)
    if ext in {"emf", "wmf"}:
        converted_path = os.path.splitext(filepath)[0] + ".jpg"
        vector_to_jpg(filepath, converted_path)
        os.remove(filepath)
        filepath = converted_path
    # path written into the HTML, relative to the HTML file's own directory
    rel_path = os.path.relpath(filepath, html_out_dir)
    return (
        f'<figure class="slide-image">'
        f'<img src="{rel_path}" alt="[TODO: caption slide {slide_idx} image {img_counter} — '
        f'not auto-generated, see script docstring]">'
        f"</figure>"
    )


# ---------- column-aware reading order ----------

def _shape_center_x(shape) -> float:
    """Return the horizontal center of a shape in EMU."""
    left = shape.left if shape.left is not None else 0
    width = shape.width if shape.width is not None else 0
    return left + width / 2


def _shape_top(shape) -> int:
    """Return the top position of a shape in EMU."""
    return shape.top if shape.top is not None else 0


def _cluster_into_columns(
    shapes: list,
    column_threshold_emu: int = 914400,  # ~1 inch in EMU
) -> list[list]:
    """
    Cluster shapes into vertical columns by horizontal centers.
    
    Shapes whose horizontal centers are within `column_threshold_emu` of each other
    are placed in the same column. Columns are then ordered left-to-right by their
    median center x. Within each column, shapes are sorted top-to-bottom.
    
    This handles simple multi-column layouts (e.g., two columns of bullets) much
    better than a pure (top, left) sort, which would interleave the columns.
    """
    if not shapes:
        return []
    
    # Compute center x for each shape
    shapes_with_center = [(s, _shape_center_x(s)) for s in shapes]
    
    # Sort by center x to form initial column groups
    shapes_with_center.sort(key=lambda x: x[1])
    
    columns = []
    current_column = []
    current_column_centers = []
    
    for shape, cx in shapes_with_center:
        if not current_column:
            current_column.append(shape)
            current_column_centers.append(cx)
        else:
            # Check if this shape belongs to the current column
            median_cx = sorted(current_column_centers)[len(current_column_centers) // 2]
            if abs(cx - median_cx) <= column_threshold_emu:
                current_column.append(shape)
                current_column_centers.append(cx)
            else:
                # Start new column
                columns.append(current_column)
                current_column = [shape]
                current_column_centers = [cx]
    
    if current_column:
        columns.append(current_column)
    
    # Sort shapes within each column by top position
    for column in columns:
        column.sort(key=_shape_top)
    
    # Columns are already left-to-right by construction
    return columns


def _get_ordered_shapes(slide) -> list:
    """
    Return shapes in reading order using column-aware clustering.
    """
    shapes = list(slide.shapes)
    columns = _cluster_into_columns(shapes)
    ordered = []
    for column in columns:
        ordered.extend(column)
    return ordered


# ---------- slide / deck assembly ----------

def _slide_to_html(
    slide,
    slide_idx: int,
    images_dir: str,
    html_out_dir: str,
) -> str:
    pieces = [f'<section class="slide" data-slide="{slide_idx}">', f"<h2>Slide {slide_idx}</h2>"]
    img_counter = 0

    shapes = _get_ordered_shapes(slide)

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img_counter += 1
            pieces.append(_picture_to_html(shape, slide_idx, img_counter, images_dir, html_out_dir))
            continue

        if getattr(shape, "has_table", False) and shape.has_table:
            pieces.append(_table_to_html(shape.table))
            continue

        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            html = _text_frame_to_html(shape.text_frame)
            if html:
                pieces.append(html)
            continue

        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            # Recurse into grouped shapes (common for diagrams built from boxes/lines)
            # Use same column-aware ordering for sub-shapes
            group_shapes = list(shape.shapes)
            group_ordered = _cluster_into_columns(group_shapes)
            for sub_group in group_ordered:
                for sub in sub_group:
                    if sub.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        img_counter += 1
                        pieces.append(_picture_to_html(sub, slide_idx, img_counter, images_dir, html_out_dir))
                    elif getattr(sub, "has_table", False) and sub.has_table:
                        pieces.append(_table_to_html(sub.table))
                    elif getattr(sub, "has_text_frame", False) and sub.has_text_frame:
                        html = _text_frame_to_html(sub.text_frame)
                        if html:
                            pieces.append(html)

    pieces.append("</section>")
    return "\n".join(pieces)


def convert(
    pptx_path: str | Path,
    html_path: str | Path,
    images_dir: Optional[str | Path] = None,
) -> tuple[Path, Path]:
    prs = Presentation(pptx_path)
    html_path = Path(html_path)
    html_out_dir = html_path.parent or Path(".")
    if images_dir is None:
        images_dir = html_out_dir / "images"
    else:
        images_dir = Path(images_dir)
    os.makedirs(images_dir, exist_ok=True)

    slides_html = []
    for i, slide in enumerate(prs.slides, start=1):
        slides_html.append(_slide_to_html(slide, i, str(images_dir), str(html_out_dir)))

    doc = f"""<!DOCTYPE html>
<html lang="fa">
<head>
<meta charset="UTF-8">
<title>{escape(os.path.basename(pptx_path))}</title>
<style>
  body {{ font-family: sans-serif; direction: rtl; }}
  mark {{ padding: 0 2px; }}
  .slide-image img {{ max-width: 500px; }}
  table {{ border-collapse: collapse; margin: 8px 0; }}
</style>
</head>
<body>
{os.linesep.join(slides_html)}
</body>
</html>
"""
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(doc)

    return html_path, images_dir


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pptx_path", type=Path)
    ap.add_argument("html_path", type=Path)
    ap.add_argument("--images-dir", type=Path, default=None, help="Defaults to <html_dir>/images")
    args = ap.parse_args()

    html_path, images_dir = convert(args.pptx_path, args.html_path, args.images_dir)
    print(f"Wrote {html_path}")
    print(f"Images in {images_dir}")


if __name__ == "__main__":
    main()
